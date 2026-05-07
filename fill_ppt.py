from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


BASE_DIR = Path("/Users/cuing/Desktop/face-lab-2315308113")
PPT_PATH = BASE_DIR / "实验一_人脸识别系统_2315308113_徐颖莎.pptx"
ARTIFACT_DIR = BASE_DIR / "artifacts"
METRICS_PATH = ARTIFACT_DIR / "metrics_summary.json"


def add_textbox(slide, left, top, width, height, text, font_size=18, bold=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.size = Pt(font_size)
    run.font.bold = bold
    return box


def add_picture(slide, image_name: str, left, top, width=None, height=None):
    path = ARTIFACT_DIR / image_name
    if width is not None and height is not None:
        slide.shapes.add_picture(str(path), left, top, width=width, height=height)
    elif width is not None:
        slide.shapes.add_picture(str(path), left, top, width=width)
    elif height is not None:
        slide.shapes.add_picture(str(path), left, top, height=height)
    else:
        slide.shapes.add_picture(str(path), left, top)


def main() -> None:
    prs = Presentation(str(PPT_PATH))
    metrics = {item["model"]: item for item in json.loads(METRICS_PATH.read_text(encoding="utf-8"))}

    slide1 = prs.slides[0]
    add_textbox(
        slide1,
        Inches(1.2),
        Inches(5.6),
        Inches(5.5),
        Inches(1.0),
        "学号：2315308113\n姓名：徐颖莎",
        font_size=22,
        bold=True,
    )

    slide2 = prs.slides[1]
    add_textbox(
        slide2,
        Inches(1.2),
        Inches(1.8),
        Inches(10.5),
        Inches(3.2),
        "实验流程：\n1. 读取 CASIA-FaceV5 数据集\n2. PCA 降维\n3. DBSCAN 去噪\n4. 训练朴素贝叶斯、SVM、随机森林、CNN\n5. 对比四种模型效果",
        font_size=22,
    )

    slide4 = prs.slides[3]
    add_textbox(
        slide4,
        Inches(0.8),
        Inches(1.5),
        Inches(4.2),
        Inches(2.0),
        "实验信息\n姓名：徐颖莎\n学号：2315308113\n数据集：CASIA-FaceV5\n最佳模型：SVM",
        font_size=20,
        bold=True,
    )
    add_picture(slide4, "metrics_table.png", Inches(5.2), Inches(1.3), width=Inches(6.6))
    add_picture(slide4, "metrics_bar.png", Inches(1.0), Inches(3.7), width=Inches(10.7))

    slide5 = prs.slides[4]
    add_picture(slide5, "dataset_samples.png", Inches(0.7), Inches(1.5), width=Inches(7.0))
    add_textbox(
        slide5,
        Inches(8.0),
        Inches(1.7),
        Inches(4.2),
        Inches(2.5),
        "数据集统计：\n- 500 名受试者\n- 2500 张灰度图像\n- 每人 5 张\n- 图像大小 128×128",
        font_size=20,
    )

    slide6 = prs.slides[5]
    add_picture(slide6, "dbscan_samples.png", Inches(0.7), Inches(1.5), width=Inches(11.0))
    add_textbox(
        slide6,
        Inches(0.9),
        Inches(5.9),
        Inches(10.8),
        Inches(0.7),
        "DBSCAN 参数：eps=24，min_samples=5；最终识别出 28 个噪声样本。",
        font_size=18,
    )

    slide7 = prs.slides[6]
    add_picture(slide7, "pca_variance.png", Inches(0.6), Inches(1.6), width=Inches(5.3))
    add_picture(slide7, "dbscan_kdist.png", Inches(6.1), Inches(1.6), width=Inches(5.3))
    add_textbox(
        slide7,
        Inches(1.0),
        Inches(5.9),
        Inches(10.6),
        Inches(0.8),
        "PCA 保留 278 维，可解释约 95% 的方差；根据 5 近邻距离曲线选择 DBSCAN 的 eps。",
        font_size=18,
    )

    slide8 = prs.slides[7]
    add_picture(slide8, "naive_bayes_cm.png", Inches(0.5), Inches(1.5), width=Inches(5.2))
    add_picture(slide8, "naive_bayes_roc.png", Inches(6.0), Inches(1.5), width=Inches(5.2))
    add_textbox(
        slide8,
        Inches(0.8),
        Inches(5.9),
        Inches(10.5),
        Inches(0.8),
        f"accuracy={metrics['naive_bayes']['accuracy']:.4f}  precision={metrics['naive_bayes']['precision']:.4f}  recall={metrics['naive_bayes']['recall']:.4f}  f1={metrics['naive_bayes']['f1']:.4f}",
        font_size=17,
    )

    slide9 = prs.slides[8]
    add_picture(slide9, "svm_cm.png", Inches(0.5), Inches(1.5), width=Inches(5.2))
    add_picture(slide9, "svm_roc.png", Inches(6.0), Inches(1.5), width=Inches(5.2))
    add_textbox(
        slide9,
        Inches(0.8),
        Inches(5.9),
        Inches(10.5),
        Inches(0.8),
        f"accuracy={metrics['svm']['accuracy']:.4f}  precision={metrics['svm']['precision']:.4f}  recall={metrics['svm']['recall']:.4f}  f1={metrics['svm']['f1']:.4f}",
        font_size=17,
    )

    slide10 = prs.slides[9]
    add_picture(slide10, "random_forest_cm.png", Inches(0.5), Inches(1.5), width=Inches(5.2))
    add_picture(slide10, "random_forest_roc.png", Inches(6.0), Inches(1.5), width=Inches(5.2))
    add_textbox(
        slide10,
        Inches(0.8),
        Inches(5.9),
        Inches(10.5),
        Inches(0.8),
        f"accuracy={metrics['random_forest']['accuracy']:.4f}  precision={metrics['random_forest']['precision']:.4f}  recall={metrics['random_forest']['recall']:.4f}  f1={metrics['random_forest']['f1']:.4f}",
        font_size=17,
    )

    slide11 = prs.slides[10]
    add_picture(slide11, "cnn_loss.png", Inches(0.4), Inches(1.4), width=Inches(3.7))
    add_picture(slide11, "cnn_cm.png", Inches(4.5), Inches(1.4), width=Inches(3.5))
    add_picture(slide11, "cnn_roc.png", Inches(8.2), Inches(1.4), width=Inches(3.5))
    add_textbox(
        slide11,
        Inches(0.8),
        Inches(5.9),
        Inches(10.5),
        Inches(0.8),
        f"accuracy={metrics['cnn']['accuracy']:.4f}  precision={metrics['cnn']['precision']:.4f}  recall={metrics['cnn']['recall']:.4f}  f1={metrics['cnn']['f1']:.4f}",
        font_size=17,
    )

    prs.save(str(PPT_PATH))


if __name__ == "__main__":
    main()
